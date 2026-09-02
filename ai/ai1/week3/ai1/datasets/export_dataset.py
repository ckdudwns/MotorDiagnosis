"""
데이터셋 매니페스트 CSV/XLSX 내보내기 (AI-1, 3주차 DATA_EXPORT_01)

register_dataset.build_manifest()가 만든 매니페스트를 받아 `output_dir/versions/<dataset
id>/`에:
1. `dataset_manifest.json` — GET /api/datasets/{id} 응답 형태(rows 제외 + artifactRefs)
2. `dataset_rows.csv`      — 학습용 정규화 행 데이터
3. `dataset_export.xlsx`   — manifest 시트 + rows 시트

를 출력하고, 그 버전 id로 `output_dir/CURRENT` 포인터를 원자적으로 전환한다(자세한 이유는
export_dataset()의 docstring 참고). week1 `build_ai1_handoff_dataset.py`처럼 CSV(csv
모듈)를 그대로 쓰고, XLSX는 요청대로 별도 무거운 도구 없이 openpyxl만 사용한다.
"""

import os
import csv
import json
import time
import shutil
import hashlib
import tempfile
import argparse

from openpyxl import Workbook, load_workbook

from register_dataset import (  # noqa: E402
    build_manifest,
    dataset_export_label_fields,
    summarize_dataset_labels,
    DATASET_EXPORT_LABEL_FIELDS,
    SPLIT_STRATEGIES,
    _DEFAULT_DATA_DIR,
    DEFAULT_SPLIT_RATIOS,
)


def _replace_with_retry(src: str, dst: str, attempts: int = 10, delay: float = 0.05) -> None:
    """os.replace()를 짧은 backoff로 재시도한다.

    Windows에서는 방금 만든 파일/디렉터리를 백신·검색 인덱서가 짧게
    핸들을 잡고 있는 동안 os.replace()가 PermissionError(WinError 5/32)로
    실패하는 경우가 흔하다 — 동시 export 테스트에서도 재현된다. 이 오류는
    금방 사라지므로 몇 차례 짧게 재시도하는 것으로 충분하고, 재시도로도
    풀리지 않으면(진짜 권한 문제 등) 마지막 예외를 그대로 전파한다.
    """
    for attempt in range(attempts):
        try:
            os.replace(src, dst)
            return
        except PermissionError:
            if attempt == attempts - 1:
                raise
            time.sleep(delay)


class VersionContentConflictError(RuntimeError):
    """같은 version_id(=dataset id)로 서로 다른 내용의 산출물을 내보내려 할 때 발생한다.

    [리뷰 P1] version_id는 `build_manifest()`의 `source.checksum`(원본 파일·전처리·
    특징 산출물 기준)으로 결정되고, export 시점에만 바뀌는 값(예: `export_dataset()`
    호출 직전에 `manifest["labelMapping"]`을 수정)은 반영하지 않는다. 그래서 같은
    id로 실제로 다른 라벨 상태를 담은 CSV를 두 번 내보내려 하면 "이미 있는
    version_dir는 방금 만든 staged 산출물과 내용이 같다"는 가정이 깨진다.
    """


def _sha256_of_file(path: str, chunk_size: int = 1 << 20) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _manifest_without_rows(manifest: dict, artifact_refs: dict, label_summary: dict) -> dict:
    """GET /api/datasets/{id} 응답 형태 (rows 제외, artifactRefs 포함).

    [리뷰 P2] labelCounts/trainingEligibleCount/trainingEligibleSplitCounts는
    manifest에 저장된(build_manifest 시점) 값이 아니라, export 시점에 rows·
    labelMapping으로 다시 계산한 `label_summary`로 덮어쓴다 — CSV/XLSX 행은
    이미 export 시점 값(dataset_export_label_fields)으로 계산하므로, 호출자가
    export 전에 manifest["labelMapping"]을 바꾸면 manifest에 캐시된 집계값과
    실제 내보낸 행이 서로 다른 라벨 상태를 보고하게 된다.
    """
    result = {k: v for k, v in manifest.items() if k != "rows"}
    result["artifactRefs"] = artifact_refs
    result["labelCounts"] = label_summary["labelCounts"]
    result["trainingEligibleCount"] = label_summary["trainingEligibleCount"]
    result["trainingEligibleSplitCounts"] = label_summary["trainingEligibleSplitCounts"]
    return result


def export_dataset(manifest: dict, output_dir: str) -> dict:
    """manifest를 CSV/XLSX/JSON으로 내보낸다.

    세 파일을 output_dir에 직접 순차 os.replace()하면 각 개별 파일 교체는
    원자적이어도 "세 파일의 집합"을 바꾸는 동작 자체는 원자적이지 않다 —
    동시에 실행된 다른 export나 그 사이 시점의 리더가 서로 다른 버전의
    CSV/XLSX/manifest가 섞인 상태를 볼 수 있다. 그래서 이 데이터셋 버전
    전용의 versioned 디렉터리(`versions/<dataset id>/`)에 세 파일을 모두
    만들고 검증한 뒤, 그 디렉터리 자체를 단 한 번의 os.replace()(디렉터리
    rename, 단일 원자적 연산)로 배치한다. 이후 `CURRENT` 포인터 파일을
    새 버전 id로 원자적으로 교체하는 것까지가 마지막 단계다 — 실패하면
    output_dir(과 CURRENT가 가리키는 버전)은 이전 상태 그대로 남는다.
    """
    os.makedirs(output_dir, exist_ok=True)
    versions_dir = os.path.join(output_dir, "versions")
    os.makedirs(versions_dir, exist_ok=True)

    rows = manifest["rows"]
    if not rows:
        raise ValueError("내보낼 행이 없습니다 (manifest['rows']가 비어 있음).")

    # 매니페스트 rows(윈도우별 특징값)는 그대로 두고, export 산출물에만 v1.3
    # DatasetExportRow 라벨 필드를 파생 컬럼으로 붙인다 — label_status/training_eligible
    # 등은 known_label + labelMapping의 결정적 함수라 특징 산출물 fingerprint에는
    # 넣지 않는다.
    label_mapping = manifest["labelMapping"]
    taxonomy_version = manifest["labelTaxonomyVersion"]
    export_rows = [
        {
            **row,
            **dataset_export_label_fields(row, label_mapping, taxonomy_version),
        }
        for row in rows
    ]
    # CSV/XLSX 행과 정확히 같은 (label_mapping, taxonomy_version)으로 다시 집계한다
    # — manifest에 캐시된 labelCounts 등을 그대로 신뢰하지 않는다(리뷰 P2).
    label_summary = summarize_dataset_labels(rows, label_mapping, taxonomy_version)

    version_id = manifest["id"]
    version_dir = os.path.join(versions_dir, version_id)

    staging_dir = tempfile.mkdtemp(prefix=".export-staging-", dir=versions_dir)
    try:
        csv_path = os.path.join(staging_dir, "dataset_rows.csv")
        fieldnames = list(rows[0].keys()) + list(DATASET_EXPORT_LABEL_FIELDS)
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(export_rows)

        xlsx_path = os.path.join(staging_dir, "dataset_export.xlsx")
        wb = Workbook()

        manifest_sheet = wb.active
        manifest_sheet.title = "manifest"
        manifest_sheet.append(["field", "value"])
        manifest_sheet.append(["id", manifest["id"]])
        manifest_sheet.append(["name", manifest["name"]])
        manifest_sheet.append(["source.type", manifest["source"]["type"]])
        manifest_sheet.append(["source.uri", manifest["source"]["uri"]])
        manifest_sheet.append(["source.license", manifest["source"]["license"]])
        manifest_sheet.append(["source.checksum", manifest["source"]["checksum"]])
        for filename, info in manifest["source"]["files"].items():
            manifest_sheet.append([f"source.files.{filename}.sha256", info["sha256"]])
            manifest_sheet.append([f"source.files.{filename}.label", info["label"]])
        manifest_sheet.append(
            ["compatibility.signalType", ",".join(manifest["compatibility"]["signalType"])]
        )
        manifest_sheet.append(
            ["compatibility.samplingRateHz", manifest["compatibility"]["samplingRateHz"]]
        )
        manifest_sheet.append(
            ["compatibility.units.vibration", manifest["compatibility"]["units"]["vibration"]]
        )
        rpm_range = manifest["compatibility"]["operatingConditions"]["rpmRange"]
        manifest_sheet.append(["compatibility.operatingConditions.rpmRange", str(rpm_range)])
        manifest_sheet.append(["labelTaxonomyVersion", manifest["labelTaxonomyVersion"]])
        manifest_sheet.append(
            ["labelPolicyVersion", manifest.get("labelPolicyVersion")]
        )
        manifest_sheet.append(
            ["snapshotSchemaVersion", manifest.get("snapshotSchemaVersion")]
        )
        for src_label, common_label in manifest["labelMapping"].items():
            manifest_sheet.append([f"labelMapping.{src_label}", common_label])
        manifest_sheet.append(["split.train", manifest["split"]["train"]])
        manifest_sheet.append(["split.validation", manifest["split"]["validation"]])
        manifest_sheet.append(["split.test", manifest["split"]["test"]])
        manifest_sheet.append(["splitStrategy", manifest["splitStrategy"]])
        manifest_sheet.append(["status", manifest["status"]])
        manifest_sheet.append(["reason", manifest["reason"]])
        manifest_sheet.append(["createdAt", manifest["createdAt"]])
        manifest_sheet.append(["rowCount", manifest["rowCount"]])
        for split_name, count in manifest["splitCounts"].items():
            manifest_sheet.append([f"splitCounts.{split_name}", count])
        for status, count in label_summary["labelCounts"].items():
            manifest_sheet.append([f"labelCounts.{status}", count])
        manifest_sheet.append(
            ["trainingEligibleCount", label_summary["trainingEligibleCount"]]
        )
        for split_name, count in label_summary["trainingEligibleSplitCounts"].items():
            manifest_sheet.append(
                [f"trainingEligibleSplitCounts.{split_name}", count]
            )

        rows_sheet = wb.create_sheet("rows")
        rows_sheet.append(fieldnames)
        for row in export_rows:
            rows_sheet.append([row.get(name) for name in fieldnames])

        wb.save(xlsx_path)

        final_csv_path = os.path.join(version_dir, "dataset_rows.csv")
        final_xlsx_path = os.path.join(version_dir, "dataset_export.xlsx")
        final_manifest_path = os.path.join(version_dir, "dataset_manifest.json")

        artifact_refs = {"csv": final_csv_path, "xlsx": final_xlsx_path}
        manifest_path = os.path.join(staging_dir, "dataset_manifest.json")
        with open(manifest_path, "w", encoding="utf-8") as f:
            json.dump(
                _manifest_without_rows(manifest, artifact_refs, label_summary),
                f,
                ensure_ascii=False,
                indent=2,
                allow_nan=False,
            )

        # 세 파일 모두 온전히 만들어졌는지 확인한 뒤에만 배치한다 — 이 중
        # 하나라도 손상돼 있으면 version_dir 자체를 만들지 않는다.
        if os.path.getsize(csv_path) == 0:
            raise RuntimeError("생성된 CSV가 비어 있습니다.")
        if os.path.getsize(xlsx_path) == 0:
            raise RuntimeError("생성된 XLSX가 비어 있습니다.")
        wb_check = load_workbook(xlsx_path)
        if "manifest" not in wb_check.sheetnames or "rows" not in wb_check.sheetnames:
            raise RuntimeError("생성된 XLSX에 manifest/rows 시트가 없습니다.")
        with open(manifest_path, encoding="utf-8") as f:
            json.load(f)  # 파싱 가능한 JSON인지 확인

        # staging_dir -> version_dir는 같은 부모(versions_dir) 안에서의
        # 디렉터리 rename이므로 단일 원자적 연산이다 — 파일별 os.replace()
        # 세 번과 달리, 중간 상태(세 파일 중 일부만 새 버전)가 관측될 수
        # 없다.
        try:
            _replace_with_retry(staging_dir, version_dir)
        except OSError:
            if not os.path.isdir(version_dir):
                raise
            # [리뷰 P1] version_id(=source.checksum)는 원본 파일·전처리·특징
            # 산출물로 결정되지만, export 시점에만 바뀌는 값(예: 호출 직전
            # manifest["labelMapping"] 수정)은 반영하지 않는다 — 그래서
            # "version_dir가 이미 있다 == 동시 export가 정확히 같은 내용을
            # 이미 배치했다"는 예전 가정이 항상 참은 아니다. CSV(rows + 라벨
            # 파생 컬럼을 그대로 담고, 타임스탬프 등 휘발성 메타데이터가 없어
            # 결정적이다 — XLSX는 openpyxl이 저장 시각을 파일에 넣어 내용이
            # 같아도 바이트가 달라지므로 비교에 쓰지 않는다)를 기존 산출물과
            # 대조해 실제로 동일한 내용인지 확인한다.
            existing_csv_path = os.path.join(version_dir, "dataset_rows.csv")
            if not os.path.exists(existing_csv_path) or (
                _sha256_of_file(csv_path) != _sha256_of_file(existing_csv_path)
            ):
                raise VersionContentConflictError(
                    f"version_id {version_id!r}가 이미 존재하지만 내보내려는 rows/라벨 "
                    "내용이 기존 산출물과 다릅니다. 이 id는 source.checksum만 반영하며 "
                    "export 시점에만 바뀐 값(예: labelMapping 변경)은 반영하지 않습니다 "
                    "— 기존 버전을 조용히 재사용하면 방금 내보내려던 내용이 사라집니다. "
                    "매니페스트를 다시 생성하거나 원인을 확인하세요."
                ) from None
            # CSV 내용이 완전히 같다 -> 동시에 실행된 다른 export가 이미 같은
            # 내용을 배치했다는 뜻이므로, 이 시도는 버리고 기존 버전을 재사용한다.
    finally:
        shutil.rmtree(staging_dir, ignore_errors=True)

    # 마지막 단계: CURRENT 포인터를 새 버전으로 원자적으로 전환한다. 여기
    # 도달했다는 것은 version_dir가 (이번에 만들었든 이미 있었든) 완전한
    # 상태라는 뜻이므로, 이 스텝이 실패해도 이전 CURRENT가 가리키는 버전은
    # 손상되지 않는다.
    current_pointer = os.path.join(output_dir, "CURRENT")
    tmp_pointer_fd, tmp_pointer_path = tempfile.mkstemp(
        prefix=".CURRENT-", dir=output_dir
    )
    try:
        with os.fdopen(tmp_pointer_fd, "w", encoding="utf-8") as f:
            f.write(version_id)
        _replace_with_retry(tmp_pointer_path, current_pointer)
    except BaseException:
        if os.path.exists(tmp_pointer_path):
            os.remove(tmp_pointer_path)
        raise

    return {
        "version_id": version_id,
        "version_dir": version_dir,
        "manifest_path": final_manifest_path,
        "csv_path": final_csv_path,
        "xlsx_path": final_xlsx_path,
        "row_count": len(rows),
    }


def resolve_current_version_dir(output_dir: str) -> str:
    """CURRENT 포인터가 가리키는 버전 디렉터리의 절대 경로를 돌려준다.

    여러 export가 동시에 실행돼도 CURRENT는 마지막에 성공적으로 전환된
    단일 버전만 가리키므로, 이 함수로 읽은 세 파일(csv/xlsx/manifest)은
    항상 같은 버전에 속한다(서로 다른 버전이 섞이지 않음).
    """
    current_pointer = os.path.join(output_dir, "CURRENT")
    with open(current_pointer, encoding="utf-8") as f:
        version_id = f.read().strip()
    return os.path.join(output_dir, "versions", version_id)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="CWRU 데이터셋 매니페스트를 CSV/XLSX로 내보내기"
    )
    parser.add_argument("--data-dir", default=_DEFAULT_DATA_DIR)
    parser.add_argument("--train-ratio", type=float, default=DEFAULT_SPLIT_RATIOS["train"])
    parser.add_argument(
        "--validation-ratio", type=float, default=DEFAULT_SPLIT_RATIOS["validation"]
    )
    parser.add_argument("--test-ratio", type=float, default=DEFAULT_SPLIT_RATIOS["test"])
    parser.add_argument(
        "--split-strategy",
        choices=SPLIT_STRATEGIES,
        default="specimen_group",
        help="specimen_group(기본, CWRU는 InsufficientAssetGroupsError) 또는 "
        "operating_condition_holdout(independentHoldout=False)",
    )
    parser.add_argument(
        "--output-dir",
        default=os.path.normpath(
            os.path.join(os.path.dirname(__file__), "..", "data", "handoff")
        ),
    )
    args = parser.parse_args()

    # 기본 specimen_group은 CWRU의 NORMAL specimen 1개 제약으로 3-way에서
    # InsufficientAssetGroupsError를 낸다 (의도된 정직한 실패).
    manifest = build_manifest(
        data_dir=args.data_dir,
        split_ratios={
            "train": args.train_ratio,
            "validation": args.validation_ratio,
            "test": args.test_ratio,
        },
        split_strategy=args.split_strategy,
    )
    result = export_dataset(manifest, args.output_dir)

    print(f"{result['row_count']}행 내보내기 완료")
    print(
        f"  라벨: {manifest['labelCounts']} / "
        f"학습가능 {manifest['trainingEligibleCount']}건 "
        f"{manifest['trainingEligibleSplitCounts']} "
        f"(정책 {manifest['labelPolicyVersion']}, 스키마 {manifest['snapshotSchemaVersion']})"
    )
    print(f"  매니페스트: {result['manifest_path']}")
    print(f"  CSV: {result['csv_path']}")
    print(f"  XLSX: {result['xlsx_path']}")
