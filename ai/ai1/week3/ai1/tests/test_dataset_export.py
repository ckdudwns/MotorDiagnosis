"""
DATA_EXPORT_01 테스트 — CWRU 데이터로 매니페스트/CSV/XLSX 내보내기 검증 (AI-1, 3주차)

실제 CWRU 데이터(ai/ai1/week1/ai1/data/external/cwru/*.mat)가 있을 때만 실행되며,
없으면 unittest.skipUnless로 명시적으로 skip 처리한다 (week2 test_week2_pipeline.py와
동일한 패턴).
"""

import os
import sys
import csv
import io
import codecs
import contextlib
import json
import shutil
import subprocess
import tempfile
import threading
import unittest
from unittest import mock

import numpy as np

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_DATASETS_DIR = os.path.normpath(os.path.join(_THIS_DIR, "..", "datasets"))
_CWRU_DATA_DIR = os.path.normpath(
    os.path.join(_THIS_DIR, "..", "..", "..", "week1", "ai1", "data", "external", "cwru")
)

sys.path.insert(0, _DATASETS_DIR)

from register_dataset import (  # noqa: E402
    build_manifest,
    group_split,
    operating_condition_split,
    compute_version_checksum,
    compute_feature_output_fingerprint,
    dataset_export_label_fields,
    summarize_dataset_labels,
    validate_split_ratios,
    sha256_of_file,
    extract_all_features,
    DATASET_LABEL_MAPPING,
    DATASET_EXPORT_LABEL_FIELDS,
    LABEL_TAXONOMY_VERSION,
    LABEL_POLICY_VERSION,
    SNAPSHOT_SCHEMA_VERSION,
    FEATURE_PIPELINE_VERSION,
    InsufficientAssetGroupsError,
)
from export_dataset import export_dataset  # noqa: E402
import extract_features as _extract_features_module  # noqa: E402
from extract_features import FeatureConfig, compute_mfcc  # noqa: E402


def _cwru_data_available() -> bool:
    return os.path.exists(os.path.join(_CWRU_DATA_DIR, "97.mat"))


CWRU_SKIP_REASON = f"CWRU 실데이터 없음: {os.path.join(_CWRU_DATA_DIR, '97.mat')}"


def _grouped_records(label: str, group_sizes: dict) -> list:
    """label에 대해 {source_label: window_count} 형태로 합성 레코드를 만든다."""
    records = []
    for source_label, count in group_sizes.items():
        records.extend(
            {"label": label, "source_label": source_label} for _ in range(count)
        )
    return records


class TestGroupSplitSynthetic(unittest.TestCase):
    """합성 레코드로 그룹 분할 로직만 검증 — 실데이터 불필요."""

    def _groups_used_per_split(self, records, splits):
        """split(train/validation/test)별로 등장한 source_label 집합을 모은다."""
        by_split: dict = {"train": set(), "validation": set(), "test": set()}
        for rec, split in zip(records, splits):
            by_split[split].add(rec["source_label"])
        return by_split

    def test_no_group_appears_in_more_than_one_split(self):
        records = _grouped_records(
            "NORMAL", {"a.mat": 20, "b.mat": 15, "c.mat": 10, "d.mat": 5}
        ) + _grouped_records(
            "BEARING_FAULT_INNER", {"e.mat": 12, "f.mat": 9, "g.mat": 6}
        )
        splits = group_split(records, seed=1)

        self.assertEqual(len(splits), len(records))
        self.assertTrue(set(splits) <= {"train", "validation", "test"})

        by_split = self._groups_used_per_split(records, splits)
        all_groups_seen = by_split["train"] | by_split["validation"] | by_split["test"]
        # 각 그룹(source_label)은 정확히 하나의 split에서만 등장해야 한다 (누수 없음)
        for group in all_groups_seen:
            memberships = sum(group in s for s in by_split.values())
            self.assertEqual(
                memberships, 1, f"그룹 {group!r}이 두 개 이상의 split에 걸쳐 있습니다."
            )

    def test_every_required_split_gets_at_least_one_group_when_enough_groups_exist(self):
        records = _grouped_records(
            "NORMAL", {"a.mat": 20, "b.mat": 15, "c.mat": 10}
        )
        splits = group_split(records, seed=1)
        by_split = self._groups_used_per_split(records, splits)
        for split_name in ("train", "validation", "test"):
            self.assertGreater(
                len(by_split[split_name]), 0, f"{split_name} split에 그룹이 배정되지 않았습니다."
            )

    def test_split_is_deterministic_given_seed(self):
        records = _grouped_records("NORMAL", {"a.mat": 10, "b.mat": 6, "c.mat": 4})
        self.assertEqual(
            group_split(records, seed=7), group_split(records, seed=7)
        )

    def test_insufficient_groups_raises_instead_of_shuffling_windows(self):
        """CWRU처럼 라벨당 그룹(자산)이 1개뿐이면, 윈도우를 섞는 대신
        명시적으로 데이터 부족 오류를 내야 한다."""
        records = _grouped_records("NORMAL", {"only.mat": 50})
        with self.assertRaises(InsufficientAssetGroupsError):
            group_split(records, seed=1)

    def test_two_groups_is_still_insufficient_for_three_way_split(self):
        records = _grouped_records("NORMAL", {"a.mat": 50, "b.mat": 50})
        with self.assertRaises(InsufficientAssetGroupsError):
            group_split(records, seed=1)

    def test_train_only_ratio_succeeds_with_a_single_group(self):
        records = _grouped_records("NORMAL", {"only.mat": 50})
        splits = group_split(
            records, ratios={"train": 1.0, "validation": 0.0, "test": 0.0}, seed=1
        )
        self.assertEqual(set(splits), {"train"})

    def test_mixed_label_group_is_not_split_across_labels(self):
        """같은 자산(source_label)이 NORMAL과 ANOMALY 레코드를 함께 가지고
        있으면, 라벨별로 그룹을 따로 만드는 예전 로직에서는 같은 자산이
        라벨에 따라 서로 다른 split에 배정될 수 있었다(예: NORMAL은 train,
        ANOMALY는 test). 자산은 라벨과 무관하게 하나의 split에만 속해야
        한다."""
        # 그룹별 크기를 라벨마다 반대 순서로 둬서, 라벨별로 그룹을 따로
        # 묶는 예전 로직이라면 shared.mat이 NORMAL 기준으로는 가장 큰 그룹(->
        # train), ANOMALY 기준으로는 가장 작은 그룹(-> test)이 되어 실제로
        # 서로 다른 split에 배정되는 상황을 재현한다.
        shared = [{"label": "NORMAL", "source_label": "shared.mat"} for _ in range(20)]
        shared += [
            {"label": "ANOMALY", "source_label": "shared.mat"} for _ in range(20)
        ]
        records = (
            shared
            + _grouped_records("NORMAL", {"b.mat": 5, "c.mat": 5})
            + _grouped_records("ANOMALY", {"e.mat": 50, "f.mat": 50})
        )
        splits = group_split(records, seed=0)

        shared_splits = {
            split
            for rec, split in zip(records, splits)
            if rec["source_label"] == "shared.mat"
        }
        self.assertEqual(
            len(shared_splits),
            1,
            f"공유 자산 shared.mat이 여러 split에 걸쳐 있습니다: {shared_splits}",
        )

    def test_split_is_independent_of_pythonhashseed(self):
        """group_split은 그룹 순회에 set을 쓰므로, 정렬 시 gid를 마지막
        tie-break로 넣지 않으면 결과가 PYTHONHASHSEED에 따라 달라진다.
        같은 records/seed로 서로 다른 PYTHONHASHSEED 하의 두 서브프로세스를
        실행해 split 배정(그리고 그로부터 나온 fingerprint)이 동일한지
        확인한다."""
        script = (
            "import sys, json\n"
            "sys.path.insert(0, %r)\n"
            "from register_dataset import group_split\n"
            "records = []\n"
            "for name, count in [('a.mat', 20), ('b.mat', 15), ('c.mat', 10), "
            "('d.mat', 5)]:\n"
            "    records += [{'label': 'NORMAL', 'source_label': name}] * count\n"
            "for name, count in [('e.mat', 12), ('f.mat', 9), ('g.mat', 6)]:\n"
            "    records += [{'label': 'BEARING_FAULT_INNER', 'source_label': name}] "
            "* count\n"
            "print(json.dumps(group_split(records, seed=1)))\n"
        ) % (_DATASETS_DIR,)

        def _run_with_hashseed(seed_value: str) -> str:
            env = dict(os.environ, PYTHONHASHSEED=seed_value)
            result = subprocess.run(
                [sys.executable, "-c", script],
                capture_output=True,
                text=True,
                env=env,
                check=True,
            )
            return result.stdout.strip()

        out_seed_1 = _run_with_hashseed("1")
        out_seed_2 = _run_with_hashseed("2")
        self.assertEqual(
            out_seed_1,
            out_seed_2,
            "PYTHONHASHSEED가 달라지면 split 배정이 달라집니다 (결정성 위반)",
        )

    def test_feasible_mixed_label_coverage_succeeds_despite_greedy_trap(self):
        """그룹이 서로 다른 라벨 쌍을 사슬처럼 공유하는 구조에서는, 각
        (라벨, split) 조합을 독립적으로 그리디 배정하면 실제로 가능한
        배정이 있어도 실패할 수 있다 — 이 테스트의 구조에서 예전 그리디는
        L3의 두 그룹(g0, g1)이 다른 라벨(L1, L0)의 요구를 채우느라 둘 다
        같은 split(test)에 몰려 배정돼, 정작 L3 자신은 train을 커버할
        그룹이 남지 않아 InsufficientAssetGroupsError를 냈다(무작위 탐색으로
        확인한 실제 반례). 최소 하나의 유효 배정이 존재하므로 예외 없이
        성공해야 한다."""
        records = (
            [{"label": "L0", "source_label": "g4.mat"}]
            + [{"label": "L0", "source_label": "g1.mat"}]
            + [{"label": "L1", "source_label": "g0.mat"}]
            + [{"label": "L1", "source_label": "g2.mat"}]
            + [{"label": "L2", "source_label": "g3.mat"}]
            + [{"label": "L2", "source_label": "g5.mat"}]
            + [{"label": "L3", "source_label": "g0.mat"}]
            + [{"label": "L3", "source_label": "g1.mat"}]
        )
        splits = group_split(
            records, ratios={"train": 0.5, "validation": 0.0, "test": 0.5}, seed=1
        )

        by_label_splits: dict = {}
        for rec, split in zip(records, splits):
            by_label_splits.setdefault(rec["label"], set()).add(split)
        for label, seen_splits in by_label_splits.items():
            self.assertEqual(
                seen_splits,
                {"train", "test"},
                f"라벨 {label!r}이 train/test를 모두 커버하지 못했습니다: {seen_splits}",
            )


class TestComputeVersionChecksum(unittest.TestCase):
    def _files(self):
        return {"97.mat": {"sha256": "sha256:aaa", "label": "NORMAL"}}

    def _checksum(self, **overrides):
        kwargs = dict(
            source_files=self._files(),
            window_size=2048,
            hop_size=2048,
            split_ratios={"train": 1.0},
            seed=42,
            label_taxonomy_version=LABEL_TAXONOMY_VERSION,
            label_mapping=DATASET_LABEL_MAPPING,
            feature_config=FeatureConfig(sample_rate=12000),
            feature_output_fingerprint="sha256:fixed-fingerprint",
            feature_pipeline_version=FEATURE_PIPELINE_VERSION,
        )
        kwargs.update(overrides)
        return compute_version_checksum(**kwargs)

    def test_deterministic_given_same_inputs(self):
        self.assertEqual(self._checksum(), self._checksum())

    def test_changes_when_window_size_changes(self):
        self.assertNotEqual(self._checksum(), self._checksum(window_size=4096))

    def test_changes_when_split_ratios_change(self):
        changed = self._checksum(split_ratios={"train": 0.5, "test": 0.5})
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_label_taxonomy_version_changes(self):
        changed = self._checksum(label_taxonomy_version="CWRU-FAULT-V2")
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_label_mapping_changes(self):
        """DATASET_LABEL_MAPPING을 재정의(예: NORMAL -> ANOMALY)해 정규화 결과가
        바뀌면, 원본 파일/window/hop/분할/seed가 같아도 체크섬이 달라져야 한다 —
        그래야 같은 id가 서로 다른 학습 데이터를 가리키는 상황을 막을 수 있다."""
        relabeled = dict(DATASET_LABEL_MAPPING)
        relabeled["NORMAL"] = "ANOMALY"
        changed = self._checksum(label_mapping=relabeled)
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_source_label_changes_with_same_sha256(self):
        """같은 sha256이라도 파일에 배정된 source label(known_label)이 바뀌면
        정규화 산출물(known_label/common_label)이 달라지므로 체크섬도 달라져야
        한다 — 이전에는 payload에서 label이 빠져 sha256만 같으면 동일했다."""
        relabeled_files = {"97.mat": {"sha256": "sha256:aaa", "label": "BEARING_FAULT_INNER"}}
        changed = self._checksum(source_files=relabeled_files)
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_feature_config_changes(self):
        """frame_length/hop_length/n_mfcc/band_edges 중 하나라도 바뀌면 특징값
        산출물이 달라지므로 체크섬도 달라져야 한다."""
        changed = self._checksum(feature_config=FeatureConfig(sample_rate=12000, n_mfcc=20))
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_feature_pipeline_version_changes(self):
        changed = self._checksum(feature_pipeline_version="week2.extract_all_features.v2")
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_feature_output_fingerprint_changes(self):
        """librosa 유무처럼 소스 코드/설정 어디에도 드러나지 않는 실행 환경
        차이(MFCC 0벡터 폴백 등)로 실제 산출된 특징값이 달라지면, 다른 입력이
        전부 같아도 fingerprint를 통해 체크섬이 달라져야 한다."""
        changed = self._checksum(feature_output_fingerprint="sha256:different-fingerprint")
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_label_policy_version_changes(self):
        """라벨 정책 버전이 바뀌면(v1.3) 같은 원본·특징이어도 다른 데이터셋 버전이어야
        한다 — 신규 데이터셋이 기존 frozen id를 재사용하는 것을 막는다."""
        changed = self._checksum(label_policy_version="LABEL-POLICY-V3")
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_snapshot_schema_version_changes(self):
        changed = self._checksum(snapshot_schema_version="3")
        self.assertNotEqual(self._checksum(), changed)

    def test_changes_when_split_strategy_changes(self):
        # specimen_group vs operating_condition_holdout는 서로 다른 데이터셋 버전이어야 한다.
        changed = self._checksum(split_strategy="operating_condition_holdout")
        self.assertNotEqual(self._checksum(), changed)


class TestComputeFeatureOutputFingerprint(unittest.TestCase):
    def test_deterministic_given_same_rows(self):
        rows = [{"sample_id": "97_0000", "mfcc_1": -151.41, "rms_mean": 0.05}]
        self.assertEqual(
            compute_feature_output_fingerprint(rows),
            compute_feature_output_fingerprint(rows),
        )

    def test_changes_when_mfcc_values_differ_like_librosa_fallback(self):
        """librosa가 있으면 실제 MFCC 값이, 없으면 0벡터가 나온다(같은 원본/설정).
        두 실행의 rows가 이 지점에서만 달라지는 상황을 재현한다."""
        rows_with_librosa = [{"sample_id": "97_0000", "mfcc_1": -151.41}]
        rows_without_librosa = [{"sample_id": "97_0000", "mfcc_1": 0.0}]
        self.assertNotEqual(
            compute_feature_output_fingerprint(rows_with_librosa),
            compute_feature_output_fingerprint(rows_without_librosa),
        )

    def test_changes_when_row_order_differs(self):
        rows_a = [{"sample_id": "a"}, {"sample_id": "b"}]
        rows_b = [{"sample_id": "b"}, {"sample_id": "a"}]
        self.assertNotEqual(
            compute_feature_output_fingerprint(rows_a),
            compute_feature_output_fingerprint(rows_b),
        )


class TestMfccFallbackWarningIsEncodingSafe(unittest.TestCase):
    """안내된 설치 방법(루트 requirements.txt + week3 requirements.txt)만으로는
    librosa가 설치되지 않는다. 이 상태로 register_dataset.py를 실행하면
    extract_all_features() -> compute_mfcc()의 librosa 미설치 폴백 경고가
    출력되는데, 이 메시지의 em dash(—)는 Windows 기본 콘솔 코드페이지인
    cp949로 표현할 수 없어 print()가 UnicodeEncodeError를 내며 rows 생성과
    fingerprint 계산 전에 실행이 통째로 중단됐다. cp949로 강제 인코딩되는
    스트림에 실제로 출력해 이 크래시가 재발하지 않는지 확인한다."""

    def test_fallback_warning_does_not_crash_on_cp949_console(self):
        cp949_stream = codecs.getwriter("cp949")(io.BytesIO())
        with mock.patch.object(_extract_features_module, "_HAS_LIBROSA", False):
            with contextlib.redirect_stdout(cp949_stream):
                mfcc = compute_mfcc(np.zeros(4096), FeatureConfig())
        self.assertTrue(np.all(mfcc == 0.0))


class TestWeek2ExtractFeaturesModuleIsolation(unittest.TestCase):
    """week1과 week2가 둘 다 `extract_features.py`라는 같은 이름의 모듈을
    갖고 있다. 같은 프로세스에서 week1의 extract_features가 먼저 평범한
    이름("extract_features")으로 import되면 sys.modules 캐시 때문에,
    이후 register_dataset이 week2 디렉터리를 sys.path 앞쪽에 넣고 같은
    이름으로 import해도 캐시된 week1 모듈이 재사용될 수 있다 — 이 경우
    week2 전용 kurtosis 특징이 조용히 빠진다."""

    def test_week2_kurtosis_feature_present_even_if_week1_module_cached_first(self):
        week1_extract_features_path = os.path.normpath(
            os.path.join(
                _THIS_DIR,
                "..",
                "..",
                "..",
                "week1",
                "ai1",
                "feature_extraction",
                "extract_features.py",
            )
        )
        import importlib.util

        spec = importlib.util.spec_from_file_location(
            "extract_features", week1_extract_features_path
        )
        week1_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(week1_module)

        # week1의 extract_features가 먼저 "extract_features"라는 이름으로
        # sys.modules에 캐시된 상황을 재현한다 (예: week1 테스트가 먼저
        # 수집·실행된 경우).
        previous = sys.modules.get("extract_features")
        sys.modules["extract_features"] = week1_module
        try:
            t = np.linspace(0, 1, 12000, endpoint=False)
            signal = np.sin(2 * np.pi * 100 * t)
            features = extract_all_features(signal, FeatureConfig(sample_rate=12000))
            self.assertIn(
                "kurtosis_mean",
                features,
                "week1 모듈 이름 충돌 때문에 week2 전용 kurtosis 특징이 빠졌습니다",
            )
        finally:
            if previous is None:
                sys.modules.pop("extract_features", None)
            else:
                sys.modules["extract_features"] = previous


class TestValidateSplitRatios(unittest.TestCase):
    def test_default_ratios_are_valid(self):
        validate_split_ratios({"train": 0.7, "validation": 0.2, "test": 0.1})

    def test_missing_key_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": 0.8, "validation": 0.2})

    def test_extra_key_rejected(self):
        """holdout처럼 실제로 쓰이지 않는 추가 키가 섞여 있으면, group_split
        결과는 정상 3-way 구성과 동일한데도 checksum payload에 그 키가 포함돼
        다른 dataset id가 생기므로 미리 거부해야 한다."""
        with self.assertRaises(ValueError):
            validate_split_ratios(
                {"train": 0.7, "validation": 0.2, "test": 0.1, "holdout": 0.0}
            )

    def test_empty_dict_rejected_not_replaced_with_default(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({})

    def test_sum_greater_than_one_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": 0.8, "validation": 0.3, "test": 0.1})

    def test_sum_less_than_one_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": 0.5, "validation": 0.2, "test": 0.1})

    def test_all_zero_rejected_instead_of_raising_max_iterable_empty(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": 0.0, "validation": 0.0, "test": 0.0})

    def test_negative_ratio_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": 1.2, "validation": -0.1, "test": -0.1})

    def test_ratio_above_one_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": 1.5, "validation": -0.5, "test": 0.0})

    def test_non_numeric_ratio_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": "0.7", "validation": 0.2, "test": 0.1})

    def test_bool_ratio_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": True, "validation": 0.0, "test": 0.0})

    def test_nan_ratio_rejected(self):
        with self.assertRaises(ValueError):
            validate_split_ratios({"train": float("nan"), "validation": 0.2, "test": 0.1})

    def test_group_split_rejects_invalid_ratios_before_max_call(self):
        """전부 0인 비율은 InsufficientAssetGroupsError(그룹 부족)가 아니라,
        group_split 내부의 max() 빈 시퀀스 오류보다 먼저 명확한 입력 검증
        오류로 걸러져야 한다."""
        records = _grouped_records("NORMAL", {"a.mat": 10, "b.mat": 10})
        with self.assertRaises(ValueError) as ctx:
            group_split(
                records, ratios={"train": 0.0, "validation": 0.0, "test": 0.0}, seed=1
            )
        self.assertNotIsInstance(ctx.exception, InsufficientAssetGroupsError)

    def test_build_manifest_rejects_invalid_ratios_before_loading_data(self):
        """CWRU 실데이터 유무와 무관하게, build_manifest는 데이터 로딩보다 먼저
        split_ratios를 검증해야 한다 (data_dir이 없어도 ValueError가 나야 함)."""
        with self.assertRaises(ValueError):
            build_manifest(
                data_dir=os.path.join(_THIS_DIR, "__no_such_dir__"),
                split_ratios={"train": 0.8, "validation": 0.3, "test": 0.1},
            )

    def test_group_split_empty_dict_rejected_not_replaced_with_default(self):
        """split_ratios={}는 `ratios or DEFAULT`의 truthiness 때문에 기본값으로
        치환되면 안 된다 — 명시적으로 빈 dict를 넘겼다면 필수 키 누락으로
        거부해야지, 조용히 기본 3-way 분할을 실행하면 안 된다."""
        records = _grouped_records("NORMAL", {"a.mat": 10, "b.mat": 10})
        with self.assertRaises(ValueError):
            group_split(records, ratios={}, seed=1)

    def test_build_manifest_empty_dict_rejected_not_replaced_with_default(self):
        with self.assertRaises(ValueError):
            build_manifest(
                data_dir=os.path.join(_THIS_DIR, "__no_such_dir__"), split_ratios={}
            )


class TestValidateFiniteFeatures(unittest.TestCase):
    def test_nan_value_rejected(self):
        from register_dataset import _validate_finite_features

        with self.assertRaises(ValueError):
            _validate_finite_features({"rms_mean": float("nan")}, "sample-1")

    def test_inf_value_rejected(self):
        from register_dataset import _validate_finite_features

        with self.assertRaises(ValueError):
            _validate_finite_features({"rms_mean": float("inf")}, "sample-1")

    def test_non_numeric_value_rejected(self):
        from register_dataset import _validate_finite_features

        with self.assertRaises(ValueError):
            _validate_finite_features({"rms_mean": "0.05"}, "sample-1")

    def test_bool_value_rejected(self):
        from register_dataset import _validate_finite_features

        with self.assertRaises(ValueError):
            _validate_finite_features({"rms_mean": True}, "sample-1")

    def test_finite_values_accepted(self):
        from register_dataset import _validate_finite_features

        _validate_finite_features({"rms_mean": 0.05, "kurtosis_mean": -1.2}, "sample-1")


@unittest.skipUnless(_cwru_data_available(), CWRU_SKIP_REASON)
class TestBuildManifestRejectsNonFiniteFeatures(unittest.TestCase):
    """extract_all_features()만 모킹하고 나머지(CWRU .mat 로딩)는 실제로
    실행하므로, 원본 파일이 없는 체크아웃에서는 build_manifest가 특징값
    검증에 도달하기 전에 FileNotFoundError로 먼저 실패한다 — 클래스
    전체를 CWRU 데이터 유무에 따라 skip한다."""

    def test_nan_feature_value_blocks_manifest_build(self):
        """NaN/Inf 특징값이 그대로 저장되면 CSV에는 문자열 "nan"이, XLSX에는
        빈 셀로 남아 같은 값이 산출물마다 다르게 표현된다 — build_manifest
        단계에서 명확히 막아야 한다."""
        with mock.patch(
            "register_dataset.extract_all_features",
            return_value={"rms_mean": float("nan")},
        ):
            with self.assertRaises(ValueError):
                build_manifest(
                    data_dir=_CWRU_DATA_DIR,
                    split_ratios={"train": 1.0, "validation": 0.0, "test": 0.0},
                )

    def test_inf_feature_value_blocks_manifest_build(self):
        with mock.patch(
            "register_dataset.extract_all_features",
            return_value={"rms_mean": float("inf")},
        ):
            with self.assertRaises(ValueError):
                build_manifest(
                    data_dir=_CWRU_DATA_DIR,
                    split_ratios={"train": 1.0, "validation": 0.0, "test": 0.0},
                )


class TestJsonDumpsRejectNonFiniteValues(unittest.TestCase):
    """저장 전 검증을 뚫고 NaN/Inf가 들어오더라도, JSON 직렬화 계층에서
    allow_nan=False가 마지막 방어선으로 명확히 실패해야 한다(기본값인
    allow_nan=True는 표준이 아닌 NaN/Infinity 리터럴을 조용히 써버린다)."""

    def test_compute_feature_output_fingerprint_rejects_nan(self):
        with self.assertRaises(ValueError):
            compute_feature_output_fingerprint(
                [{"sample_id": "x", "rms_mean": float("nan")}]
            )

    def test_compute_feature_output_fingerprint_rejects_inf(self):
        with self.assertRaises(ValueError):
            compute_feature_output_fingerprint(
                [{"sample_id": "x", "rms_mean": float("inf")}]
            )


class TestOperatingConditionSplit(unittest.TestCase):
    """부하조건 기준 고정 배정 — 결정적(seed 무관), 부하 tier 통째로 한 split에.
    같은 물리 베어링이 여러 부하에 걸쳐 있으므로 specimen 독립 검증이 아니다."""

    def _recs(self, n_per_hp=3):
        recs = []
        for specimen in ("S-A", "S-B"):
            for hp in (0, 1, 2, 3):
                recs += [
                    {"specimen_id": specimen, "load_hp": hp, "label": "NORMAL"}
                ] * n_per_hp
        return recs

    def test_deterministic_load_tier_assignment(self):
        recs = self._recs()
        a = operating_condition_split(recs)
        b = operating_condition_split(recs, ratios={"train": 0.1})  # ratios 무시
        self.assertEqual(a, b)
        # 0HP->test, 1HP->validation, 2·3HP->train
        for rec, split in zip(recs, a):
            expected = {0: "test", 1: "validation", 2: "train", 3: "train"}[rec["load_hp"]]
            self.assertEqual(split, expected)

    def test_same_specimen_spans_multiple_splits(self):
        recs = self._recs()
        splits = operating_condition_split(recs)
        by_specimen = {}
        for rec, split in zip(recs, splits):
            by_specimen.setdefault(rec["specimen_id"], set()).add(split)
        self.assertTrue(all(len(s) == 3 for s in by_specimen.values()))

    def test_missing_load_hp_raises(self):
        with self.assertRaises(ValueError):
            operating_condition_split([{"specimen_id": "x", "label": "NORMAL"}])


class TestDatasetExportLabelFields(unittest.TestCase):
    """API 명세서 v1.3 DatasetExportRow: CWRU known_label은 신뢰된 외부 라벨이므로
    매핑 성공 행은 verified/training_eligible, 매핑 실패는 unmapped, 라벨 없음은
    unlabeled. scenario_label 등 텔레메트리 전용 필드는 항상 None."""

    _TAX = "CWRU-FAULT-V1"

    def test_mapped_known_label_is_verified_and_trainable(self):
        info = dataset_export_label_fields(
            {"known_label": "BEARING_FAULT_INNER", "common_label": "ANOMALY"},
            DATASET_LABEL_MAPPING, self._TAX,
        )
        self.assertEqual(info["label_status"], "verified")
        self.assertTrue(info["training_eligible"])
        self.assertEqual(info["ground_truth_label"], "BEARING_FAULT_INNER")
        self.assertEqual(info["ground_truth_source"], "dataset_registration")
        self.assertEqual(info["target_label"], "ANOMALY")
        self.assertEqual(info["target_label_taxonomy_version"], self._TAX)
        self.assertIsNone(info["scenario_label"])
        self.assertFalse(info["event_reviewed"])

    def test_unmapped_known_label_is_excluded_from_training(self):
        info = dataset_export_label_fields(
            {"known_label": "UNSEEN_FAULT"}, DATASET_LABEL_MAPPING, self._TAX
        )
        self.assertEqual(info["label_status"], "unmapped")
        self.assertFalse(info["training_eligible"])
        self.assertEqual(info["ground_truth_label"], "UNSEEN_FAULT")
        self.assertIsNone(info["target_label"])
        self.assertIsNone(info["target_label_taxonomy_version"])

    def test_missing_known_label_is_unlabeled(self):
        for row in ({"known_label": None}, {"known_label": ""}, {}):
            info = dataset_export_label_fields(row, DATASET_LABEL_MAPPING, self._TAX)
            self.assertEqual(info["label_status"], "unlabeled")
            self.assertFalse(info["training_eligible"])
            self.assertIsNone(info["ground_truth_label"])
            self.assertIsNone(info["ground_truth_source"])

    def test_summarize_counts_and_split_breakdown(self):
        rows = [
            {"known_label": "NORMAL", "split": "train"},
            {"known_label": "BEARING_FAULT_BALL", "split": "validation"},
            {"known_label": "UNSEEN", "split": "test"},
            {"known_label": None, "split": "train"},
        ]
        summary = summarize_dataset_labels(rows, DATASET_LABEL_MAPPING, self._TAX)
        self.assertEqual(
            summary["labelCounts"],
            {"verified": 2, "weak": 0, "unlabeled": 1, "unmapped": 1},
        )
        self.assertEqual(summary["trainingEligibleCount"], 2)
        self.assertEqual(
            summary["trainingEligibleSplitCounts"],
            {"train": 1, "validation": 1, "test": 0},
        )


class TestExportDatasetSynthetic(unittest.TestCase):
    """CWRU 실데이터 없이도 openpyxl XLSX 내보내기 자체를 검증하는 합성 매니페스트 테스트."""

    def _synthetic_manifest(self):
        rows = [
            {
                "sample_id": "97_0000",
                "source_file": "97.mat",
                "known_label": "NORMAL",
                "common_label": "NORMAL",
                "split": "train",
                "sample_rate_hz": 12000,
                "rpm": 1797,
                "rms_mean": 0.05,
            },
            {
                "sample_id": "105_0000",
                "source_file": "105.mat",
                "known_label": "BEARING_FAULT_INNER",
                "common_label": "ANOMALY",
                "split": "test",
                "sample_rate_hz": 12000,
                "rpm": 1797,
                "rms_mean": 0.42,
            },
        ]
        return {
            "id": "DS-CWRU-VIBRATION-19700101-deadbeefcafe",
            "name": "cwru-bearing-vibration-v1",
            "source": {
                "type": "external",
                "uri": "https://example.invalid/cwru",
                "license": "test",
                "checksum": "sha256:deadbeef",
                "files": {"97.mat": {"sha256": "sha256:aaa", "label": "NORMAL"}},
            },
            "compatibility": {
                "signalType": ["vibration"],
                "samplingRateHz": 12000,
                "units": {"vibration": "g"},
                "operatingConditions": {"rpmRange": [1797, 1797], "load": "test"},
            },
            "labelTaxonomyVersion": "CWRU-FAULT-V1",
            "labelMapping": DATASET_LABEL_MAPPING,
            "labelPolicyVersion": LABEL_POLICY_VERSION,
            "snapshotSchemaVersion": SNAPSHOT_SCHEMA_VERSION,
            "split": {"train": 0.5, "validation": 0.0, "test": 0.5},
            "splitStrategy": "test",
            "status": "draft",
            "reason": "synthetic test",
            "createdAt": "1970-01-01T00:00:00+00:00",
            "rowCount": len(rows),
            "splitCounts": {"train": 1, "validation": 0, "test": 1},
            **summarize_dataset_labels(rows, DATASET_LABEL_MAPPING, "CWRU-FAULT-V1"),
            "rows": rows,
        }

    def test_export_produces_xlsx_with_manifest_and_rows_sheets(self):
        from openpyxl import load_workbook

        manifest = self._synthetic_manifest()
        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_synthetic_")
        try:
            result = export_dataset(manifest, tmp_dir)
            self.assertTrue(os.path.exists(result["xlsx_path"]))

            wb = load_workbook(result["xlsx_path"])
            self.assertIn("manifest", wb.sheetnames)
            self.assertIn("rows", wb.sheetnames)
            self.assertEqual(wb["rows"].max_row - 1, manifest["rowCount"])
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_export_csv_has_v13_label_columns_with_verified_status(self):
        manifest = self._synthetic_manifest()
        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_labels_")
        try:
            result = export_dataset(manifest, tmp_dir)
            with open(result["csv_path"], newline="", encoding="utf-8") as f:
                csv_rows = list(csv.DictReader(f))
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

        for field in DATASET_EXPORT_LABEL_FIELDS:
            self.assertIn(field, csv_rows[0])
        # 합성 매니페스트의 두 known_label은 모두 labelMapping에 있다 → verified.
        for row in csv_rows:
            self.assertEqual(row["label_status"], "verified")
            self.assertEqual(row["training_eligible"], "True")
            self.assertEqual(row["ground_truth_source"], "dataset_registration")
            self.assertEqual(row["ground_truth_label"], row["known_label"])
            self.assertEqual(row["target_label"], row["common_label"])
            self.assertEqual(row["target_label_taxonomy_version"], "CWRU-FAULT-V1")
            self.assertEqual(row["scenario_label"], "")  # CSV 빈 셀
            self.assertEqual(row["event_reviewed"], "False")

    def test_export_manifest_json_has_label_summary_and_policy_versions(self):
        manifest = self._synthetic_manifest()
        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_summary_")
        try:
            result = export_dataset(manifest, tmp_dir)
            with open(result["manifest_path"], encoding="utf-8") as f:
                exported = json.load(f)
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

        self.assertEqual(exported["labelPolicyVersion"], LABEL_POLICY_VERSION)
        self.assertEqual(exported["snapshotSchemaVersion"], SNAPSHOT_SCHEMA_VERSION)
        self.assertEqual(sum(exported["labelCounts"].values()), exported["rowCount"])
        self.assertEqual(exported["labelCounts"]["verified"], exported["rowCount"])
        self.assertEqual(exported["trainingEligibleCount"], exported["rowCount"])
        self.assertEqual(
            sum(exported["trainingEligibleSplitCounts"].values()),
            exported["trainingEligibleCount"],
        )

    def test_exported_manifest_label_summary_matches_csv_not_stale_cache(self):
        """[리뷰 P2] JSON manifest의 labelCounts/trainingEligibleCount는
        build_manifest 시점에 캐시된 값이 아니라 export 시점 (rows, labelMapping)
        으로 다시 계산해야 한다. labelMapping을 export 직전에 바꾸면 CSV/XLSX
        행은 새 매핑을 반영하는데 JSON은 예전 집계값을 그대로 보고하던 문제를
        고정한다."""
        manifest = self._synthetic_manifest()
        # BEARING_FAULT_INNER 매핑을 빼서 두 번째 행(105_0000)이 unmapped가
        # 되도록 한다 — labelMapping은 새 dict로 교체(공유 전역 객체를
        # 직접 변형하지 않음).
        manifest["labelMapping"] = {"NORMAL": "NORMAL"}
        # manifest["labelCounts"]/trainingEligibleCount는 (의도적으로) 예전
        # 매핑 기준의 오래된 값 그대로 둔다 — export_dataset이 이를 신뢰하지
        # 않아야 한다.
        self.assertEqual(manifest["labelCounts"]["verified"], 2)  # 오래된(잘못된) 캐시

        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_stale_labels_")
        try:
            result = export_dataset(manifest, tmp_dir)
            with open(result["csv_path"], newline="", encoding="utf-8") as f:
                csv_rows = {row["sample_id"]: row for row in csv.DictReader(f)}
            with open(result["manifest_path"], encoding="utf-8") as f:
                exported = json.load(f)
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

        self.assertEqual(csv_rows["97_0000"]["label_status"], "verified")
        self.assertEqual(csv_rows["105_0000"]["label_status"], "unmapped")
        self.assertEqual(csv_rows["105_0000"]["training_eligible"], "False")

        # JSON은 CSV와 같은(export 시점 재계산) 집계를 보고해야 한다 — 오래된
        # "verified: 2" 캐시가 아니라.
        self.assertEqual(exported["labelCounts"]["verified"], 1)
        self.assertEqual(exported["labelCounts"]["unmapped"], 1)
        self.assertEqual(exported["trainingEligibleCount"], 1)

    def test_failed_export_does_not_corrupt_previous_version(self):
        """CSV/XLSX/manifest를 output_dir에 바로 순차 기록하면, 뒤쪽 파일
        생성이 실패했을 때 앞서 이미 덮어쓴 파일만 새 버전이고 나머지는
        이전 버전으로 남아 서로 다른 버전이 섞인다. XLSX 저장이 실패해도
        v1 버전 디렉터리와 CURRENT 포인터가 실패 이전 상태 그대로여야 한다."""
        manifest_v1 = self._synthetic_manifest()
        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_atomic_")
        try:
            result_v1 = export_dataset(manifest_v1, tmp_dir)
            with open(result_v1["csv_path"], "rb") as f:
                csv_before = f.read()
            with open(result_v1["xlsx_path"], "rb") as f:
                xlsx_before = f.read()
            with open(result_v1["manifest_path"], "rb") as f:
                manifest_before = f.read()
            current_before = os.path.join(tmp_dir, "CURRENT")
            with open(current_before, encoding="utf-8") as f:
                current_contents_before = f.read()

            manifest_v2 = self._synthetic_manifest()
            manifest_v2["id"] = "DS-CWRU-VIBRATION-19700102-deadbeefcafe"
            manifest_v2["rows"][0]["rms_mean"] = 0.99  # v1과 구분되는 값

            with mock.patch(
                "export_dataset.Workbook.save", side_effect=RuntimeError("disk full")
            ):
                with self.assertRaises(RuntimeError):
                    export_dataset(manifest_v2, tmp_dir)

            with open(result_v1["csv_path"], "rb") as f:
                self.assertEqual(f.read(), csv_before, "실패한 내보내기가 CSV를 건드렸습니다")
            with open(result_v1["xlsx_path"], "rb") as f:
                self.assertEqual(f.read(), xlsx_before, "실패한 내보내기가 XLSX를 건드렸습니다")
            with open(result_v1["manifest_path"], "rb") as f:
                self.assertEqual(
                    f.read(), manifest_before, "실패한 내보내기가 manifest를 건드렸습니다"
                )
            with open(current_before, encoding="utf-8") as f:
                self.assertEqual(
                    f.read(),
                    current_contents_before,
                    "실패한 내보내기가 CURRENT 포인터를 건드렸습니다",
                )

            v2_version_dir = os.path.join(tmp_dir, "versions", manifest_v2["id"])
            self.assertFalse(
                os.path.exists(v2_version_dir),
                "실패한 내보내기가 v2 버전 디렉터리를 만들어 남겼습니다",
            )

            versions_dir = os.path.join(tmp_dir, "versions")
            leftovers = [
                name
                for name in os.listdir(versions_dir)
                if name.startswith(".export-staging-")
            ]
            self.assertEqual(
                leftovers, [], f"스테이징 디렉터리가 정리되지 않았습니다: {leftovers}"
            )
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_export_writes_files_under_versioned_directory_with_current_pointer(self):
        from export_dataset import resolve_current_version_dir

        manifest = self._synthetic_manifest()
        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_versioned_")
        try:
            result = export_dataset(manifest, tmp_dir)
            self.assertEqual(result["version_id"], manifest["id"])
            self.assertEqual(result["version_dir"], resolve_current_version_dir(tmp_dir))
            self.assertTrue(result["csv_path"].startswith(result["version_dir"]))
            self.assertTrue(os.path.exists(os.path.join(tmp_dir, "CURRENT")))
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_repeated_export_of_identical_manifest_is_idempotent(self):
        """version_id는 산출물 내용으로 결정되는 불변 체크섬을 담고 있으므로,
        같은 매니페스트를 두 번 내보내면 같은 version_dir을 재사용해야
        한다(디렉터리가 이미 있다는 이유로 실패하면 안 됨)."""
        manifest = self._synthetic_manifest()
        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_idempotent_")
        try:
            result_a = export_dataset(manifest, tmp_dir)
            result_b = export_dataset(manifest, tmp_dir)
            self.assertEqual(result_a["version_dir"], result_b["version_dir"])
            with open(result_b["csv_path"], "rb") as f:
                self.assertGreater(len(f.read()), 0)
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_concurrent_exports_of_different_manifests_never_expose_mixed_version(self):
        """서로 다른 두 버전을 동시에 내보내도, CURRENT가 가리키는 버전의
        csv/xlsx/manifest 세 파일은 항상 같은 버전에서 나온 것이어야 한다
        (한쪽 버전의 CSV와 다른 쪽 버전의 XLSX가 섞여 보이면 안 됨)."""
        from export_dataset import resolve_current_version_dir

        manifest_a = self._synthetic_manifest()
        manifest_a["id"] = "DS-CWRU-VIBRATION-CONCURRENT-A"
        manifest_a["rows"][0]["rms_mean"] = 0.11

        manifest_b = self._synthetic_manifest()
        manifest_b["id"] = "DS-CWRU-VIBRATION-CONCURRENT-B"
        manifest_b["rows"][0]["rms_mean"] = 0.22

        tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_concurrent_")
        try:
            errors = []

            def _export(manifest):
                try:
                    export_dataset(manifest, tmp_dir)
                except Exception as exc:  # pragma: no cover - 실패 시 진단용
                    errors.append(exc)

            for _ in range(20):
                t_a = threading.Thread(target=_export, args=(manifest_a,))
                t_b = threading.Thread(target=_export, args=(manifest_b,))
                t_a.start()
                t_b.start()
                t_a.join()
                t_b.join()

                self.assertEqual(errors, [])

                current_version_dir = resolve_current_version_dir(tmp_dir)
                self.assertIn(
                    os.path.basename(current_version_dir),
                    (manifest_a["id"], manifest_b["id"]),
                )
                with open(
                    os.path.join(current_version_dir, "dataset_manifest.json"),
                    encoding="utf-8",
                ) as f:
                    exported_manifest = json.load(f)
                with open(
                    os.path.join(current_version_dir, "dataset_rows.csv"),
                    newline="",
                    encoding="utf-8",
                ) as f:
                    exported_rows = list(csv.DictReader(f))

                # CURRENT가 가리키는 버전의 manifest.id와 실제로 그 디렉터리에
                # 있는 CSV의 rms_mean이 항상 같은 버전 쌍(A-A 또는 B-B)이어야
                # 한다 — 서로 다른 버전의 파일이 섞여 있으면 안 된다.
                if exported_manifest["id"] == manifest_a["id"]:
                    self.assertEqual(exported_rows[0]["rms_mean"], "0.11")
                else:
                    self.assertEqual(exported_rows[0]["rms_mean"], "0.22")
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)


@unittest.skipUnless(_cwru_data_available(), CWRU_SKIP_REASON)
class TestSpecimenGroupingForCwru(unittest.TestCase):
    """[리뷰 P1] CWRU 부하별 파일은 독립 자산이 아니라 같은 물리 베어링(specimen)이다.
    group split은 specimen 단위로 해야 하고, NORMAL specimen은 1개뿐이라 기본
    3-way는 정직하게 실패해야 한다. 데모용 operating_condition_holdout은 비독립임을
    플래그로 표시한다."""

    def test_specimen_group_3way_fails_because_normal_has_one_specimen(self):
        with self.assertRaises(InsufficientAssetGroupsError):
            build_manifest(data_dir=_CWRU_DATA_DIR, seed=42)  # 기본 specimen_group 3-way

    def test_fault_only_records_do_support_specimen_independent_3way(self):
        # NORMAL을 뺀 결함 3클래스는 각 specimen 3개(0.007/0.014/0.021")라 진짜
        # specimen 독립 3-way split이 가능하다 — grouping 자체가 동작함을 증명.
        # (register_dataset import 시 load_cwru_vibration 경로가 sys.path에 추가됨)
        from load_cwru_vibration import load_cwru_dataset as _load

        records = [
            r for r in _load(_CWRU_DATA_DIR) if r["label"] != "NORMAL"
        ]
        splits = group_split(records, seed=42, group_key="specimen_id")
        by_specimen = {}
        for rec, split in zip(records, splits):
            by_specimen.setdefault(rec["specimen_id"], set()).add(split)
        for specimen, s in by_specimen.items():
            self.assertEqual(len(s), 1, f"{specimen}가 여러 split에: {s}")
        self.assertEqual(set(splits), {"train", "validation", "test"})

    def test_operating_condition_holdout_manifest_split_matches_actual_rows(self):
        """[리뷰 P1] operating_condition_split은 요청 split_ratios를 완전히 무시하고
        부하 tier로 배정을 고정한다. 예전 build_manifest는 이 무시된 요청값을 그대로
        매니페스트["split"]/체크섬에 기록해서, 예를 들어 train=1.0/validation=
        test=0.0을 넘겨도 실제 rows에는 validation/test가 들어가는데 매니페스트는
        "전부 train"이라고 거짓말했다. 기록된 split은 실제 splitCounts 비율과
        일치해야 한다."""
        manifest = build_manifest(
            data_dir=_CWRU_DATA_DIR,
            split_strategy="operating_condition_holdout",
            split_ratios={"train": 1.0, "validation": 0.0, "test": 0.0},
        )
        total = manifest["rowCount"]
        for name in ("train", "validation", "test"):
            expected_ratio = manifest["splitCounts"][name] / total
            self.assertAlmostEqual(manifest["split"][name], expected_ratio, places=9)
        # 요청한 대로 "전부 train"이 되지는 않았다 — validation/test가 실제로 존재.
        self.assertGreater(manifest["splitCounts"]["validation"], 0)
        self.assertGreater(manifest["splitCounts"]["test"], 0)

    def test_operating_condition_holdout_succeeds_but_is_not_independent(self):
        manifest = build_manifest(
            data_dir=_CWRU_DATA_DIR, split_strategy="operating_condition_holdout"
        )
        self.assertFalse(manifest["independentHoldout"])
        self.assertEqual(manifest["holdoutType"], "operating_condition")
        self.assertEqual(sum(manifest["splitCounts"].values()), manifest["rowCount"])
        self.assertTrue(all(c > 0 for c in manifest["splitCounts"].values()))
        # 같은 물리 베어링이 train/validation/test에 함께 들어간다 (비독립).
        splits_by_specimen = {}
        for row in manifest["rows"]:
            splits_by_specimen.setdefault(row["specimen_id"], set()).add(row["split"])
        self.assertTrue(
            any(len(s) > 1 for s in splits_by_specimen.values()),
            "operating_condition_holdout인데 specimen이 split을 안 넘나든다?",
        )


@unittest.skipUnless(_cwru_data_available(), CWRU_SKIP_REASON)
class TestRegisterAndExportRealCwruData(unittest.TestCase):
    """실제 CWRU 데이터로 매니페스트 생성 → CSV/XLSX 내보내기까지 전체 흐름을 검증.

    기본 specimen_group은 NORMAL specimen 1개 제약으로 실패하므로
    operating_condition_holdout(비독립, independentHoldout=False)로 파이프라인
    전체(분할/체크섬/라벨매핑/내보내기)를 검증한다.
    """

    @classmethod
    def setUpClass(cls):
        cls.manifest = build_manifest(
            data_dir=_CWRU_DATA_DIR, split_strategy="operating_condition_holdout"
        )
        cls.tmp_dir = tempfile.mkdtemp(prefix="ai1_week3_dataset_export_")
        cls.export_result = export_dataset(cls.manifest, cls.tmp_dir)

    def test_holdout_is_flagged_non_independent(self):
        self.assertFalse(self.manifest["independentHoldout"])
        self.assertEqual(self.manifest["holdoutType"], "operating_condition")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp_dir, ignore_errors=True)

    def test_row_count_matches_manifest(self):
        self.assertEqual(self.manifest["rowCount"], len(self.manifest["rows"]))
        self.assertGreater(self.manifest["rowCount"], 0)

    def test_split_counts_sum_to_row_count(self):
        total = sum(self.manifest["splitCounts"].values())
        self.assertEqual(total, self.manifest["rowCount"])

    def test_checksums_match_recomputed_values(self):
        for filename, info in self.manifest["source"]["files"].items():
            recomputed = sha256_of_file(os.path.join(_CWRU_DATA_DIR, filename))
            self.assertEqual(info["sha256"], recomputed)

    def test_manifest_has_feature_output_fingerprint(self):
        self.assertIn("featureOutputFingerprint", self.manifest)
        self.assertTrue(self.manifest["featureOutputFingerprint"].startswith("sha256:"))

    def test_operating_condition_holdout_checksum_independent_of_unused_seed(self):
        """[리뷰 P2] operating_condition_split은 seed를 전혀 쓰지 않는다. 실제
        rows/split이 동일하면 seed만 바꿔도 checksum/id가 달라지면 안 된다."""
        other = build_manifest(
            data_dir=_CWRU_DATA_DIR,
            seed=999,
            split_strategy="operating_condition_holdout",
        )
        self.assertEqual(self.manifest["source"]["checksum"], other["source"]["checksum"])
        self.assertEqual(self.manifest["id"], other["id"])

    def test_id_and_source_checksum_change_with_real_config_change(self):
        # window_size는 operating_condition_holdout에서도 실제로 rows/특징을
        # 바꾸는 설정이므로(체크섬 payload에 포함) 다른 버전 체크섬/ID가 나와야 한다.
        other = build_manifest(
            data_dir=_CWRU_DATA_DIR,
            window_size=1024,
            hop_size=1024,
            seed=42,
            split_strategy="operating_condition_holdout",
        )
        self.assertIn("checksum", self.manifest["source"])
        self.assertNotEqual(self.manifest["source"]["checksum"], other["source"]["checksum"])
        self.assertNotEqual(self.manifest["id"], other["id"])
        self.assertIn(self.manifest["source"]["checksum"].split(":", 1)[1][:12], self.manifest["id"])

    def test_id_and_checksum_ignore_requested_split_ratios(self):
        """[리뷰 P1] operating_condition_split은 요청 split_ratios를 완전히
        무시한다. 실제 데이터가 같으면(요청 비율만 다르게 줘도) 매니페스트
        id/checksum이 같아야 한다 — 예전에는 무시된 입력이 체크섬 payload에
        그대로 들어가 실제로 동일한 데이터가 다른 버전으로 판정됐다."""
        other = build_manifest(
            data_dir=_CWRU_DATA_DIR,
            split_ratios={"train": 1.0, "validation": 0.0, "test": 0.0},
            seed=42,
            split_strategy="operating_condition_holdout",
        )
        self.assertEqual(self.manifest["source"]["checksum"], other["source"]["checksum"])
        self.assertEqual(self.manifest["id"], other["id"])
        self.assertEqual(self.manifest["split"], other["split"])

    def test_label_mapping_applied_to_every_row(self):
        for row in self.manifest["rows"]:
            expected = DATASET_LABEL_MAPPING[row["known_label"]]
            self.assertEqual(row["common_label"], expected)

    def test_exported_csv_row_count_matches_manifest(self):
        with open(self.export_result["csv_path"], newline="", encoding="utf-8") as f:
            csv_rows = list(csv.DictReader(f))
        self.assertEqual(len(csv_rows), self.manifest["rowCount"])

    def test_all_cwru_rows_verified_and_training_eligible(self):
        # CWRU 원본 라벨 4종은 모두 labelMapping에 있으므로 전 행 verified.
        self.assertEqual(
            self.manifest["labelCounts"],
            {
                "verified": self.manifest["rowCount"],
                "weak": 0,
                "unlabeled": 0,
                "unmapped": 0,
            },
        )
        self.assertEqual(
            self.manifest["trainingEligibleCount"], self.manifest["rowCount"]
        )
        self.assertEqual(
            sum(self.manifest["trainingEligibleSplitCounts"].values()),
            self.manifest["rowCount"],
        )
        self.assertEqual(self.manifest["labelPolicyVersion"], LABEL_POLICY_VERSION)
        self.assertEqual(
            self.manifest["snapshotSchemaVersion"], SNAPSHOT_SCHEMA_VERSION
        )

    def test_exported_csv_carries_v13_label_columns(self):
        with open(self.export_result["csv_path"], newline="", encoding="utf-8") as f:
            csv_rows = list(csv.DictReader(f))
        for field in DATASET_EXPORT_LABEL_FIELDS:
            self.assertIn(field, csv_rows[0])
        self.assertTrue(all(r["label_status"] == "verified" for r in csv_rows))
        self.assertTrue(all(r["training_eligible"] == "True" for r in csv_rows))

    def test_exported_manifest_json_excludes_rows_and_has_artifact_refs(self):
        with open(self.export_result["manifest_path"], encoding="utf-8") as f:
            exported = json.load(f)
        self.assertNotIn("rows", exported)
        self.assertIn("artifactRefs", exported)
        self.assertTrue(os.path.exists(exported["artifactRefs"]["csv"]))
        self.assertTrue(os.path.exists(exported["artifactRefs"]["xlsx"]))

    def test_xlsx_has_manifest_and_rows_sheets(self):
        from openpyxl import load_workbook

        wb = load_workbook(self.export_result["xlsx_path"])
        self.assertIn("manifest", wb.sheetnames)
        self.assertIn("rows", wb.sheetnames)
        rows_sheet = wb["rows"]
        # 헤더 1행 + 데이터 행 수
        self.assertEqual(rows_sheet.max_row - 1, self.manifest["rowCount"])


if __name__ == "__main__":
    unittest.main()
